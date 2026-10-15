"""
report_sections/front_matter.py — Cover page, Title page, Declaration, Certificate,
Abstract, Acknowledgements, and Preliminaries conforming to VIT Chennai guidelines.
"""

import docx
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from report_sections.doc_helpers import add_p


def generate_front_matter(doc: docx.Document):
    """Generate all preliminary pages matching the exact VIT template layout."""

    # ═════════════════════════════════════════════════════════════════════
    # 1. COVER PAGE
    # ═════════════════════════════════════════════════════════════════════
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(36)
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run("A project report on")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)
    r.italic = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(24)
    r = p.add_run("AI-POWERED DRONE SYSTEM FOR REAL-TIME CROWD MANAGEMENT AND MONITORING")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(20)
    r.bold = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run("Submitted in partial fulfillment for the award of the degree of")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)
    r.italic = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(24)
    r = p.add_run("Bachelor of Technology in Computer Science and Engineering")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(20)
    r.bold = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run("by")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)
    r.italic = True

    # Template Spec: 1ST & 2ND PAGE LIST ALL THE TEAM MEMBERS
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(36)
    r = p.add_run("ROHAN DUBASI (23BRS1119)\nTEJAS RAMESH BODETTI (23BRS1180)")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(16)
    r.bold = True

    # Logo placeholder / Callout
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(20)
    p.paragraph_format.space_after = Pt(20)
    r = p.add_run("[INSERT VIT CHENNAI EMBLEM / LOGO HERE]")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)
    r.bold = True
    r.font.color.rgb = RGBColor(120, 120, 120)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run("SCHOOL OF COMPUTER SCIENCE AND ENGINEERING\nVELLORE INSTITUTE OF TECHNOLOGY, CHENNAI")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(16)
    r.bold = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("November, 2025")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)

    doc.add_page_break()

    # ═════════════════════════════════════════════════════════════════════
    # 2. TITLE PAGE (Identical text to Cover Page)
    # ═════════════════════════════════════════════════════════════════════
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(36)
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run("A project report on")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)
    r.italic = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(24)
    r = p.add_run("AI-POWERED DRONE SYSTEM FOR REAL-TIME CROWD MANAGEMENT AND MONITORING")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(20)
    r.bold = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run("Submitted in partial fulfillment for the award of the degree of")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)
    r.italic = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(24)
    r = p.add_run("Bachelor of Technology in Computer Science and Engineering")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(20)
    r.bold = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run("by")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)
    r.italic = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(36)
    r = p.add_run("ROHAN DUBASI (23BRS1119)\nTEJAS RAMESH BODETTI (23BRS1180)")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(16)
    r.bold = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(20)
    p.paragraph_format.space_after = Pt(20)
    r = p.add_run("[INSERT VIT CHENNAI EMBLEM / LOGO HERE]")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)
    r.bold = True
    r.font.color.rgb = RGBColor(120, 120, 120)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run("SCHOOL OF COMPUTER SCIENCE AND ENGINEERING\nVELLORE INSTITUTE OF TECHNOLOGY, CHENNAI")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(16)
    r.bold = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("November, 2025")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)

    doc.add_page_break()

    # ═════════════════════════════════════════════════════════════════════
    # 3. DECLARATION (Only Candidate's Name)
    # ═════════════════════════════════════════════════════════════════════
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(20)
    p.paragraph_format.space_after = Pt(20)
    r = p.add_run("[INSERT VIT LOGO]")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11)
    r.bold = True
    r.font.color.rgb = RGBColor(120, 120, 120)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(24)
    r = p.add_run("DECLARATION")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)
    r.bold = True
    r.underline = True

    add_p(doc,
          "I hereby declare that the thesis entitled “AI-POWERED DRONE SYSTEM FOR REAL-TIME CROWD MANAGEMENT AND MONITORING” "
          "submitted by TEJAS RAMESH BODETTI (23BRS1180), for the award of the degree of Bachelor of Technology in "
          "Computer Science and Engineering, Vellore Institute of Technology, Chennai is a record of bonafide work carried out "
          "by me under the supervision of Prof. Valarmathi Sudhakar.")

    add_p(doc,
          "I further declare that the work reported in this thesis has not been submitted and will not be submitted, "
          "either in part or in full, for the award of any other degree or diploma in this institute or any other institute or university.")

    p_sig = doc.add_paragraph()
    p_sig.paragraph_format.space_before = Pt(60)
    p_sig.paragraph_format.line_spacing = 1.3
    r = p_sig.add_run("Place: Chennai\nDate:  [Date]\n\n\n\n")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)
    r.bold = True

    p_right = doc.add_paragraph()
    p_right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_cand = p_right.add_run("Signature of the Candidate\nTEJAS RAMESH BODETTI (23BRS1180)")
    r_cand.font.name = 'Times New Roman'
    r_cand.font.size = Pt(12)
    r_cand.bold = True

    doc.add_page_break()

    # ═════════════════════════════════════════════════════════════════════
    # 4. CERTIFICATE
    # ═════════════════════════════════════════════════════════════════════
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(16)
    r = p.add_run("[INSERT VIT LOGO]")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11)
    r.bold = True
    r.font.color.rgb = RGBColor(120, 120, 120)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run("School of Computer Science and Engineering")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)
    r.bold = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run("CERTIFICATE")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)
    r.bold = True
    r.underline = True

    add_p(doc,
          "This is to certify that the report entitled “AI-POWERED DRONE SYSTEM FOR REAL-TIME CROWD MANAGEMENT AND MONITORING” "
          "is prepared and submitted by TEJAS RAMESH BODETTI (23BRS1180) to Vellore Institute of Technology, Chennai, in partial "
          "fulfillment of the requirement for the award of the degree of Bachelor of Technology in Computer Science and Engineering "
          "is a bonafide record carried out under my guidance. The project fulfills the requirements as per the regulations of this "
          "University and in my opinion meets the necessary standards for submission. The contents of this report have not been "
          "submitted and will not be submitted either in part or in full, for the award of any other degree or diploma and the same is certified.")

    p_guide = doc.add_paragraph()
    p_guide.paragraph_format.space_before = Pt(36)
    p_guide.paragraph_format.space_after = Pt(24)
    r = p_guide.add_run("Signature of the Guide: _________________________\n"
                        "Name: Prof. Valarmathi Sudhakar\n"
                        "Date: [Date]\n")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11.5)

    p_exam = doc.add_paragraph()
    p_exam.paragraph_format.space_before = Pt(16)
    p_exam.paragraph_format.space_after = Pt(24)
    r = p_exam.add_run("Signature of the Examiner 1: ________________      Signature of the Examiner 2: ________________\n"
                       "Name: [Examiner 1 Name]                         Name: [Examiner 2 Name]\n"
                       "Date: [Date]                                    Date: [Date]\n")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11)

    p_hod = doc.add_paragraph()
    p_hod.paragraph_format.space_before = Pt(16)
    p_hod.paragraph_format.space_after = Pt(0)
    r = p_hod.add_run("Approved by the Head of Department,\n"
                      "Department of Computer Science and Engineering\n\n"
                      "Name: [HoD Name]\n"
                      "Date: [Date]\n"
                      "                                                     (Seal of the School)\n")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11.5)

    doc.add_page_break()

    # ═════════════════════════════════════════════════════════════════════
    # 5. ABSTRACT
    # ═════════════════════════════════════════════════════════════════════
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(24)
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run("ABSTRACT")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)
    r.bold = True
    r.underline = True

    add_p(doc,
          "Large public gatherings, cultural festivals, sports events, and religious congregations frequently encounter severe "
          "stampede risks, localized overcrowding, and catastrophic pedestrian bottlenecks. Conventional static closed-circuit "
          "television (CCTV) infrastructure is severely restricted by fixed fields of view, blind spots, and lack of active intervention "
          "mechanisms. This thesis presents the design, mathematical modeling, aerodynamic payload sizing, and implementation of a "
          "distributed, ground-assisted robotic drone system for real-time aerial crowd monitoring and acoustic intervention.")

    add_p(doc,
          "The mechanical and aeronautical design centers on a multi-rotor quadcopter platform engineered with an emphasis on strict payload "
          "aerodynamics, disk loading, and electrical power budgets. The drone carries two specialized payloads: an ultra-high-definition "
          "GoPro Hero 13 optical acquisition system (190 g with custom vibration-isolated mount) and an ESP32-based crowd management acoustic "
          "intervention unit (130 g) featuring a high-efficiency horn speaker driven by a Class-D amplifier through built-in digital-to-analog "
          "conversion (DAC GPIO 25) with an active-LOW emergency trigger input (GPIO 14). Total All-Up Weight (AUW) is mathematically derived "
          "at 1.68 kg, necessitating a comprehensive Blade Element Momentum Theory (BEMT) and momentum formulation to ensure a thrust-to-weight "
          "ratio (TWR) exceeding 2.1:1 using 2216 880KV brushless motors and 1045 propellers powered by a 4S 5200 mAh Lithium-Polymer battery, "
          "yielding an endurance of 18.05 minutes.")

    add_p(doc,
          "To eliminate the prohibitive computational mass and power penalties of onboard neural network acceleration, a ground-assisted "
          "distributed computing paradigm is established. Video streaming latency is critically minimized (~100 ms) via a custom dual-adapter "
          "networking topology where Wi-Fi Adapter 1 establishes a direct UDP stream to the GoPro Access Point (10.5.5.9), while Wi-Fi Adapter 2 "
          "bridges ground-station telemetry, client dashboard WebSockets, and a wireless TCP socket server (port 8002) dedicated to the ESP32 speaker. "
          "The ground-station vision pipeline executes single-stage YOLOv8 person detection paired with ByteTrack multi-object tracking. By diagnosing "
          "and rectifying track ID churn via confidence thresholding (0.25) and Kalman filter association buffers, identity stability increased by 52.5% "
          "across dense aerial scenes (~43 people). A 3×4 spatial grid evaluates regional density, maps spatial centroids to geographic coordinates via "
          "pinhole camera projection and ArduPilot state estimation, and triggers automated acoustic signals (sirens, attention chimes, and multilingual "
          "text-to-speech dispersal notices) alongside an interactive dark-themed operations dashboard with live demo footage switching.")

    add_p(doc,
          "Keywords: Aerial Robotics, Quadcopter Payload Dynamics, Momentum Theory, Crowd Monitoring, Dual-Adapter Telemetry, "
          "GoPro Low-Latency UDP, ESP32 DAC Audio, YOLOv8, ByteTrack, Grid Density Estimation, Pinhole Georeferencing.",
          bold=True)

    doc.add_page_break()

    # ═════════════════════════════════════════════════════════════════════
    # 6. ACKNOWLEDGEMENT (Individual submission by Tejas Ramesh Bodetti)
    # ═════════════════════════════════════════════════════════════════════
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(24)
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run("ACKNOWLEDGEMENT")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)
    r.bold = True
    r.underline = True

    add_p(doc,
          "It is my pleasure to express with a deep sense of gratitude to Prof. Valarmathi Sudhakar, "
          "School of Computer Science and Engineering, Vellore Institute of Technology, Chennai, for her constant guidance, "
          "continual encouragement, understanding; more than all, she taught me patience in my endeavor. My association with her "
          "is not confined to academics only, but it is a great opportunity on my part of work with an intellectual and expert "
          "in the field of autonomous systems and computer science.")

    add_p(doc,
          "It is with gratitude that I would like to extend my thanks to the visionary leader Dr. G. Viswanathan, our Honorable Chancellor, "
          "Mr. Sankar Viswanathan, Dr. Sekar Viswanathan, Dr. G. V. Selvam, Vice Presidents, Dr. Sandhya Pentareddy, Executive Director, "
          "Ms. Kadhambari S. Viswanathan, Assistant Vice-President, Dr. V. S. Kanchana Bhaaskaran, Vice-Chancellor, Dr. T. Thyagarajan, "
          "Pro-Vice Chancellor, VIT Chennai, and Dr. P. K. Manoharan, Additional Registrar, for providing an exceptional working environment "
          "and inspiring all of us during the tenure of the course.")

    add_p(doc,
          "Special mention to Dr. Viswanathan V, Dean, Dr. Nithyanandam P, Associate Dean, Dr. Suganya G, Associate Dean, Dr. Sweetlin Hemalatha C, "
          "Associate Dean, School of Computer Science and Engineering, Vellore Institute of Technology, Chennai for spending their valuable time "
          "and efforts in sharing their knowledge and for helping us in every aspect.")

    add_p(doc,
          "In jubilant state, I express ingeniously my whole-hearted thanks to [HoD Name], Head of the Department, "
          "B.Tech. Computer Science and Engineering and the Project Coordinators for their valuable support and encouragement "
          "to take up and complete the thesis.")

    add_p(doc,
          "My sincere thanks to all the faculty and staff at Vellore Institute of Technology, Chennai who helped me acquire the requisite "
          "knowledge. I would like to thank my parents for their unconditional support, sacrifice, and encouragement throughout my academic journey. "
          "It is indeed a pleasure to thank my team member, Rohan Dubasi, and my friends who encouraged me to take up and complete this task.")

    p_sig = doc.add_paragraph()
    p_sig.paragraph_format.space_before = Pt(40)
    r = p_sig.add_run("Place: Chennai\nDate:  [Date]\n\n\nTEJAS RAMESH BODETTI (23BRS1180)")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)
    r.bold = True

    doc.add_page_break()

    # ═════════════════════════════════════════════════════════════════════
    # 7. TABLE OF CONTENTS
    # ═════════════════════════════════════════════════════════════════════
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(24)
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run("TABLE OF CONTENTS")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)
    r.bold = True

    toc_items = [
        ("DECLARATION", "iii"),
        ("CERTIFICATE", "iv"),
        ("ABSTRACT", "v"),
        ("ACKNOWLEDGEMENT", "vi"),
        ("LIST OF FIGURES", "ix"),
        ("LIST OF TABLES", "xi"),
        ("LIST OF SYMBOLS AND ABBREVIATIONS", "xii"),
        ("CHAPTER 1: INTRODUCTION", "1"),
        ("    1.1 Overview and Background Context", "1"),
        ("    1.2 Motivation: Aerial Robotics in Crowd Disasters", "3"),
        ("    1.3 Problem Statement", "5"),
        ("    1.4 Objectives of the Project", "6"),
        ("    1.5 Scope of the Project (In-Scope & Out-of-Scope)", "7"),
        ("    1.6 Feasibility and Technical Risk Analysis", "9"),
        ("    1.7 Ethical Governance, Privacy and Regulatory Compliance", "11"),
        ("    1.8 Work Breakdown and Individual Contributions", "13"),
        ("    1.9 Organization of the Thesis Report", "15"),
        ("CHAPTER 2: LITERATURE SURVEY AND BACKGROUND", "17"),
        ("    2.1 Evolution of Multi-Rotor UAVs in Surveillance", "17"),
        ("    2.2 Multi-Rotor Aerodynamics and Thrust Physics", "19"),
        ("    2.3 Acoustic and Loudspeaker Payloads for Drone Crowd Control", "22"),
        ("    2.4 Wireless Video Streaming, Latency Optimization and Protocols", "25"),
        ("    2.5 Edge vs. Ground-Station Distributed Processing Paradigms", "28"),
        ("    2.6 Computer Vision Models for Aerial Pedestrian Detection", "31"),
        ("    2.7 Multi-Object Tracking Algorithms in High-Density Scenes", "34"),
        ("    2.8 Spatial Crowd Density Estimation and Flow Modeling", "37"),
        ("    2.9 Critical Review of Existing Systems and Patents", "40"),
        ("    2.10 Identification of Research Gaps", "42"),
        ("CHAPTER 3: SYSTEM REQUIREMENTS AND MATHEMATICAL MODELLING OF DRONE ROBOTICS & PAYLOAD", "44"),
        ("    3.1 Aerodynamic Momentum Theory Formulation for Multi-Rotors", "44"),
        ("    3.2 Blade Element Momentum Theory (BEMT) and Inflow Mechanics", "48"),
        ("    3.3 Multi-Rotor Platform Topology: Quadcopter vs. Hexacopter vs. Octocopter", "52"),
        ("    3.4 Airframe Mechanical Evaluation: F450 vs. S500 vs. Tarot 650", "55"),
        ("    3.5 Detailed Payload Mass & Dimensional Breakdown", "58"),
        ("        3.5.1 Primary Optical Payload: GoPro Hero 13 Black", "58"),
        ("        3.5.2 Secondary Acoustic Payload: ESP32 Crowd Management Unit", "62"),
        ("        3.5.3 Complete Itemized All-Up Weight (AUW) Mass Budget", "65"),
        ("    3.6 Propulsion Sizing and Thrust-to-Weight Ratio (TWR) Analysis", "67"),
        ("        3.6.1 Dynamic Thrust Requirements and Safety Margins", "67"),
        ("        3.6.2 Brushless DC Motor Sizing: 2212 920KV vs. 2216 880KV vs. 2312 960KV", "70"),
        ("        3.6.3 Propeller Selection and Aerodynamic Disk Loading (1045 vs. 9450)", "73"),
        ("        3.6.4 Electronic Speed Controller (ESC) Sizing and Thermal Headroom", "76"),
        ("    3.7 Electrical Power Budget and Hover Endurance Mathematical Derivation", "78"),
        ("        3.7.1 Lithium-Polymer Battery Chemistry and Usable Energy Sizing", "78"),
        ("        3.7.2 Mathematical Derivation of Hover Flight Time", "81"),
        ("        3.7.3 Electrical Current Draw Across Flight Regimes", "84"),
        ("    3.8 Flight Dynamics, Center of Gravity (CG) Shift and Inertia Tensor", "86"),
        ("        3.8.1 Three-Dimensional Center of Gravity Calculation", "86"),
        ("        3.8.2 Mass Moment of Inertia Tensor (Ixx, Iyy, Izz) Formulation", "89"),
        ("        3.8.3 Underslung Payload Pendulum Oscillations and Dynamic Stability", "92"),
        ("        3.8.4 Aerodynamic Drag Estimation and Wind Tilt Angle", "95"),
        ("    3.9 Mechanical Vibration Dampening and Acoustic Isolation Analysis", "97"),
        ("CHAPTER 4: GROUND-ASSISTED ARCHITECTURE & WIRELESS COMMUNICATION INFRASTRUCTURE", "101"),
        ("    4.1 Distributed Ground-Assisted Architectural Topology", "101"),
        ("    4.2 Onboard Flight Control Architecture (Pixhawk 2.4.8 & ArduPilot)", "104"),
        ("    4.3 Dual-Adapter High-Performance Wireless Networking Architecture", "108"),
        ("    4.4 Low-Latency Video Pipeline Implementation (gopro_client.py)", "112"),
        ("    4.5 ESP32 Wireless Audio Streaming & Button Interrupt Protocol (esp32_manager.py)", "116"),
        ("    4.6 WebSocket Telemetry and Asynchronous State Synchronization", "121"),
        ("CHAPTER 5: COMPUTER VISION, DENSITY ESTIMATION & REAL-TIME ALERT GENERATION", "124"),
        ("    5.1 Object Detection Architecture Comparison: YOLOv8 Variants", "124"),
        ("    5.2 Multi-Object Tracking with ByteTrack and ID Churn Mitigation", "130"),
        ("    5.3 3×4 Spatial Grid Density Partitioning and Density Classification", "136"),
        ("    5.4 Georeferencing and Pinhole Camera GPS Zone Localization", "140"),
        ("    5.5 Multi-Tier Rule-Based Alert Decision Engine", "144"),
        ("CHAPTER 6: SYSTEM INTEGRATION, HARDWARE ASSEMBLY & MISSION DASHBOARD", "148"),
        ("    6.1 Drone Airframe and Avionics Hardware Assembly", "148"),
        ("    6.2 GoPro Hero 13 Mechanical Mounting and Isolation", "151"),
        ("    6.3 ESP32 Crowd Speaker Hardware Fabrication & Wiring", "154"),
        ("    6.4 Ground Station Software Server Pipeline (analysis_server.py)", "157"),
        ("    6.5 Mission Control Web Dashboard Implementation", "161"),
        ("    6.6 Seamless Video Source Toggling & Offline Demo Loop Engine", "166"),
        ("CHAPTER 7: EXPERIMENTAL RESULTS AND PERFORMANCE EVALUATION", "170"),
        ("    7.1 Propulsion Dynamometer and Thrust Benchmarking Results", "170"),
        ("    7.2 Flight Hover Endurance Experimental Measurements", "174"),
        ("    7.3 Video Latency and Wireless Throughput Evaluation", "177"),
        ("    7.4 Detection, Tracking, and Density Estimation Evaluation", "180"),
        ("    7.5 ESP32 Acoustic Sound Pressure Level (SPL) and Range Testing", "185"),
        ("    7.6 End-to-End System Latency and Emergency Response Pipeline", "188"),
        ("CHAPTER 8: CONCLUSION AND FUTURE WORK", "191"),
        ("    8.1 Summary of Contributions", "191"),
        ("    8.2 Limitations and Critical Insights", "193"),
        ("    8.3 Future Work and Roadmap", "195"),
        ("REFERENCES (APA FORMAT)", "198"),
        ("APPENDIX 1: ESP32 FIRMWARE SOURCE CODE (crowd_speaker.ino)", "205"),
        ("APPENDIX 2: GROUND STATION SERVER CONFIGURATION & LAUNCH SCRIPTS", "211"),
        ("APPENDIX 3: MOTOR & PROPELLER THRUST DYNAMOMETER DATA SHEETS", "215"),
    ]

    p_tbl = doc.add_table(rows=len(toc_items), cols=2)
    p_tbl.autofit = False
    for i, (title_str, page_str) in enumerate(toc_items):
        c0, c1 = p_tbl.rows[i].cells
        c0.width = docx.shared.Inches(5.7)
        c1.width = docx.shared.Inches(0.8)

        p0 = c0.paragraphs[0]
        p0.paragraph_format.line_spacing = 1.15
        p0.paragraph_format.space_after = Pt(2)
        r0 = p0.add_run(title_str)
        r0.font.name = 'Times New Roman'
        r0.font.size = Pt(11)
        if "CHAPTER" in title_str or title_str.isupper():
            r0.bold = True

        p1 = c1.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p1.paragraph_format.line_spacing = 1.15
        p1.paragraph_format.space_after = Pt(2)
        r1 = p1.add_run(page_str)
        r1.font.name = 'Times New Roman'
        r1.font.size = Pt(11)
        if "CHAPTER" in title_str or title_str.isupper():
            r1.bold = True

    doc.add_page_break()

    # ═════════════════════════════════════════════════════════════════════
    # 8. LIST OF FIGURES & LIST OF TABLES & ABBREVIATIONS
    # ═════════════════════════════════════════════════════════════════════
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(24)
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run("LIST OF FIGURES")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)
    r.bold = True

    figures_list = [
        ("Figure 1.1", "System Architecture Overview: Aerial Subsystem, Ground Station, and Dashboard", "3"),
        ("Figure 1.2", "Project Work Breakdown and Execution Timeline", "14"),
        ("Figure 2.1", "Actuator Disk Momentum Theory Control Volume and Streamtube", "20"),
        ("Figure 2.2", "Blade Element Aerodynamics: Relative Inflow Velocity and Sectional Forces", "21"),
        ("Figure 3.1", "Structural Comparison of Evaluated Multi-Rotor Airframes (F450 vs. S500 vs. Tarot 650)", "56"),
        ("Figure 3.2", "Fully Assembled S500 Quadcopter Drone with Pixhawk 2.4.8 and FlySky Radio System", "57"),
        ("Figure 3.3", "GoPro Hero 13 Underslung Camera Mounting Bracket and Vibration Isolation System", "61"),
        ("Figure 3.4", "ESP32 Crowd Management Acoustic & Visual Unit Circuit Schematic and Speaker Assembly", "64"),
        ("Figure 3.5", "Motor Thrust vs. Throttle Percentage Curves for 2212 920KV and 2216 880KV Motors", "72"),
        ("Figure 3.6", "Propeller Thrust and Mechanical Power Curves as a Function of Rotational Speed (RPM)", "75"),
        ("Figure 3.7", "Battery State of Charge (SoC) vs. Terminal Voltage Discharge Curve Under 4S Load", "80"),
        ("Figure 3.8", "Hover Flight Endurance as a Function of All-Up Weight (AUW) and Battery Capacity", "83"),
        ("Figure 3.9", "Three-Dimensional Center of Gravity Coordinate Frame and Principal Axes of Inertia", "88"),
        ("Figure 3.10", "Transmissibility Ratio Curve of Silicone Damping Isolators Across Vibration Frequencies", "99"),
        ("Figure 4.1", "Distributed Ground-Assisted Robotic System Architecture and Data Flow", "103"),
        ("Figure 4.2", "Pixhawk 2.4.8 Internal Sensor Bus Architecture and Fail-safe Redundancy Loops", "107"),
        ("Figure 4.3", "Dual-Adapter High-Performance Wireless Network Topology on Ground Station", "110"),
        ("Figure 4.4", "Low-Latency UDP Video Ingestion and MJPEG Streaming Pipeline Architecture", "115"),
        ("Figure 4.5", "ESP32 Wireless Audio Streaming & Button Interrupt Protocol Timing Diagram", "119"),
        ("Figure 4.6", "WebSocket Telemetry and Asynchronous State Synchronization Message Topology", "123"),
        ("Figure 5.1", "Ultralytics YOLOv8 Single-Stage Neural Network Architecture (CSPDarknet + PANet)", "127"),
        ("Figure 5.2", "Person Detection Visual Output on Aerial Crowd Test Clip (42–45 People Detected)", "129"),
        ("Figure 5.3", "ByteTrack Two-Pass Matching Pipeline and Kalman Filter State Flowchart", "133"),
        ("Figure 5.4", "Persistent Multi-Object Tracking Bounding Boxes with Individual IDs", "134"),
        ("Figure 5.5", "Track ID Churn Reduction: Comparison Before (conf=0.10) and After (conf=0.25 + Buffer)", "135"),
        ("Figure 5.6", "3×4 Spatial Grid Partitioning Overlaid on Video Frame with Per-Zone Occupancy Counts", "138"),
        ("Figure 5.7", "Pinhole Camera Coordinate Geometry for Image-Space to GPS Ground Projection", "143"),
        ("Figure 5.8", "Multi-Tier Rule-Based Alert Decision State Machine", "146"),
        ("Figure 6.1", "Complete Drone Avionics and Propulsion Power Wiring Diagram", "150"),
        ("Figure 6.2", "GoPro Hero 13 Underslung Quick-Release Buckle Installation on Drone Frame", "153"),
        ("Figure 6.3", "ESP32 Acoustic Unit Hardware Wiring Breadboard and PCB Layout", "156"),
        ("Figure 6.4", "Ground Station Software Architecture and Inter-Process Communication", "160"),
        ("Figure 6.5", "Mission Control Web Dashboard Interface (Dark-Themed Operations Center UI)", "163"),
        ("Figure 6.6", "Interactive ESP32 Crowd Speaker Control Panel and Multilingual Voice Broadcast Widget", "165"),
        ("Figure 6.7", "Seamless Video Source Toggling Controls (GoPro Live vs. Demo Sample Clip vs. Webcam)", "168"),
        ("Figure 7.1", "Measured Motor Thrust Curves and Electrical Efficiency (gf/W) Benchmarking", "172"),
        ("Figure 7.2", "Experimental Flight Hover Endurance vs. Added Payload Mass", "176"),
        ("Figure 7.3", "Video Transmission Latency Profile across UDP Direct Link vs. Standard Wi-Fi Hotspot", "179"),
        ("Figure 7.4", "Frame-by-Frame People Detection Count Stability across 341 Video Frames", "183"),
        ("Figure 7.5", "Acoustic Sound Pressure Level (SPL) Decay Curve Over Outdoor Distance (dB vs. Meters)", "187"),
        ("Figure 7.6", "End-to-End System Latency Breakdown from Video Capture to Audio Dispersal Signal", "190"),
    ]

    for num, caption, pg in figures_list:
        p_fig = doc.add_paragraph()
        p_fig.paragraph_format.line_spacing = 1.15
        p_fig.paragraph_format.space_after = Pt(2)
        r0 = p_fig.add_run(f"{num}: ")
        r0.font.name = 'Times New Roman'
        r0.font.size = Pt(11)
        r0.bold = True
        r1 = p_fig.add_run(f"{caption} ")
        r1.font.name = 'Times New Roman'
        r1.font.size = Pt(11)
        r2 = p_fig.add_run(f"........................................................................ {pg}")
        r2.font.name = 'Times New Roman'
        r2.font.size = Pt(10)

    doc.add_page_break()

    # List of Tables
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(24)
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run("LIST OF TABLES")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)
    r.bold = True

    tables_list = [
        ("Table 2.1", "Comparative Analysis of State-of-the-Art UAV Crowd Monitoring Systems", "27"),
        ("Table 2.2", "Summary of Object Detection Architectures for Small-Scale Aerial Targets", "33"),
        ("Table 3.1", "Structural Comparison of Multi-Rotor Airframe Configurations", "54"),
        ("Table 3.2", "Multi-Rotor Topology Evaluation Matrix (Quadcopter vs. Hexacopter vs. Octocopter)", "55"),
        ("Table 3.3", "GoPro Hero 13 Specifications, Optical Parameters, and Power Consumption", "60"),
        ("Table 3.4", "ESP32 Crowd Management Acoustic Unit Component Mass & Power Breakdown", "63"),
        ("Table 3.5", "Itemized Mass Budget and Weight Distribution (AUW Breakdown)", "66"),
        ("Table 3.6", "Brushless DC Motor Sizing & Dynamometer Benchmarking at 14.8V (4S)", "71"),
        ("Table 3.7", "Propeller Aerodynamic Performance Matching (1045 vs. 9450 vs. 1147)", "74"),
        ("Table 3.8", "Electronic Speed Controller (ESC) Specification and Thermal Headroom", "77"),
        ("Table 3.9", "Lithium-Polymer Battery Capacity vs. Discharge Profiles", "79"),
        ("Table 3.10", "Subsystem Electrical Power Demands Across Operating Flight Regimes", "85"),
        ("Table 3.11", "Three-Dimensional Center of Gravity (CG) Coordinates and Inertial Tensor Matrix", "91"),
        ("Table 3.12", "Transmissibility Ratios of Silicone Damper Mounts Across Vibration Frequencies", "98"),
        ("Table 4.1", "Avionics Sensor Suite Specifications for Pixhawk 2.4.8", "106"),
        ("Table 4.2", "Dual-Adapter Network Interface Architecture & IP Subnet Allocation", "111"),
        ("Table 4.3", "Video Streaming Transport Protocol Comparison (UDP vs. RTSP vs. WebRTC vs. RTMP)", "113"),
        ("Table 4.4", "ESP32 Binary Audio Streaming Protocol Frame Definition", "118"),
        ("Table 5.1", "Ultralytics YOLOv8 Architecture Performance Comparison (Nano to X-Large)", "128"),
        ("Table 5.2", "ByteTrack Two-Pass Tracking Hyperparameters and Filter Tuning Values", "132"),
        ("Table 5.3", "Spatial Grid 3×4 Coordinate Bounds and Density Classification Bands", "139"),
        ("Table 5.4", "Pinhole Camera Coordinate Frame Matrix Parameters", "142"),
        ("Table 5.5", "Multi-Tier Rule-Based Alert Trigger Logic Matrix", "145"),
        ("Table 6.1", "Electronic Circuit Pin Connections for ESP32 Crowd Speaker Payload", "155"),
        ("Table 6.2", "Ground Station Software Modules and Execution Roles", "159"),
        ("Table 7.1", "Measured Motor Thrust and Electrical Efficiency across Throttle Settings", "171"),
        ("Table 7.2", "Flight Endurance Experimental Measurements with Varying Payload Masses", "175"),
        ("Table 7.3", "Network Streaming Latency and Packet Drop Benchmarking", "178"),
        ("Table 7.4", "YOLOv8x Detection Metrics on Aerial Test Clip (341 Frames)", "182"),
        ("Table 7.5", "ByteTrack Tracking Consistency and ID Churn Reduction Metrics", "184"),
        ("Table 7.6", "Acoustic Sound Pressure Level (SPL) Attenuation vs. Distance", "186"),
        ("Table 7.7", "End-to-End System Response Times for Crowd Dispersion Workflow", "189"),
    ]

    for num, caption, pg in tables_list:
        p_tbl_line = doc.add_paragraph()
        p_tbl_line.paragraph_format.line_spacing = 1.15
        p_tbl_line.paragraph_format.space_after = Pt(2)
        r0 = p_tbl_line.add_run(f"{num}: ")
        r0.font.name = 'Times New Roman'
        r0.font.size = Pt(11)
        r0.bold = True
        r1 = p_tbl_line.add_run(f"{caption} ")
        r1.font.name = 'Times New Roman'
        r1.font.size = Pt(11)
        r2 = p_tbl_line.add_run(f"........................................................................ {pg}")
        r2.font.name = 'Times New Roman'
        r2.font.size = Pt(10)

    doc.add_page_break()

    # List of Symbols, Abbreviations and Nomenclature
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(24)
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run("LIST OF SYMBOLS, ABBREVIATIONS AND NOMENCLATURE")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)
    r.bold = True

    abbreviations = [
        ("UAV", "Unmanned Aerial Vehicle"),
        ("AUW", "All-Up Weight (Total takeoff mass of the drone including all payloads)"),
        ("TWR", "Thrust-to-Weight Ratio"),
        ("BEMT", "Blade Element Momentum Theory"),
        ("ESC", "Electronic Speed Controller"),
        ("BLDC", "Brushless Direct Current (Electric Motor)"),
        ("LiPo", "Lithium-Polymer (Secondary rechargeable battery chemistry)"),
        ("DoD", "Depth of Discharge"),
        ("CG", "Center of Gravity"),
        ("IMU", "Inertial Measurement Unit (Gyroscopes and Accelerometers)"),
        ("GPS", "Global Positioning System"),
        ("ToF", "Time-of-Flight (Optical distance measurement)"),
        ("LiDAR", "Light Detection and Ranging"),
        ("DAC", "Digital-to-Analog Converter"),
        ("PCM", "Pulse-Code Modulation (Uncompressed digital audio format)"),
        ("SPL", "Sound Pressure Level (Measured in decibels, dB)"),
        ("AP", "Access Point (Wi-Fi network mode)"),
        ("STA", "Station Mode (Wi-Fi client mode)"),
        ("UDP", "User Datagram Protocol"),
        ("TCP", "Transmission Control Protocol"),
        ("MJPEG", "Motion Joint Photographic Experts Group (Video stream format)"),
        ("RTSP", "Real-Time Streaming Protocol"),
        ("TTS", "Text-to-Speech (Automated voice synthesis)"),
        ("YOLO", "You Only Look Once (Single-stage convolutional object detection)"),
        ("IoU", "Intersection over Union (Bounding box overlap metric)"),
        ("MOT", "Multi-Object Tracking"),
        ("mAP", "Mean Average Precision"),
        ("FPS", "Frames Per Second"),
        ("GSD", "Ground Sampling Distance (Real-world size of one image pixel)"),
        ("FOV", "Field of View (Angular lens coverage)"),
        ("NMS", "Non-Maximum Suppression (Bounding box filtering technique)"),
        ("PANet", "Path Aggregation Network (Feature pyramid fusion architecture)"),
        ("CSP", "Cross-Stage Partial (Neural network backbone architecture)"),
        ("MAVLink", "Micro Air Vehicle Link (Standardized lightweight drone telemetry protocol)"),
        ("DDS", "Data Distribution Service (Real-time publish-subscribe middleware)"),
        ("ROS", "Robot Operating System"),
    ]

    abbrev_tbl = doc.add_table(rows=len(abbreviations), cols=2)
    abbrev_tbl.autofit = False
    for i, (short, full) in enumerate(abbreviations):
        c0, c1 = abbrev_tbl.rows[i].cells
        c0.width = docx.shared.Inches(1.8)
        c1.width = docx.shared.Inches(4.7)

        p0 = c0.paragraphs[0]
        p0.paragraph_format.line_spacing = 1.15
        p0.paragraph_format.space_after = Pt(2)
        r0 = p0.add_run(short)
        r0.font.name = 'Times New Roman'
        r0.font.size = Pt(11)
        r0.bold = True

        p1 = c1.paragraphs[0]
        p1.paragraph_format.line_spacing = 1.15
        p1.paragraph_format.space_after = Pt(2)
        r1 = p1.add_run(full)
        r1.font.name = 'Times New Roman'
        r1.font.size = Pt(11)

    doc.add_page_break()
