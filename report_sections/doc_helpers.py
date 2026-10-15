"""
report_sections/doc_helpers.py — Document formatting and styling engine
Adheres strictly to the Vellore Institute of Technology (VIT) Chennai Project Report Template.
"""

import docx
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls


def create_base_document() -> docx.Document:
    """Initialize Document with A4 dimensions and VIT Chennai margin specs."""
    doc = docx.Document()
    section = doc.sections[0]
    section.page_width = Inches(8.27)   # A4 Width: 21.0 cm
    section.page_height = Inches(11.69) # A4 Height: 29.7 cm

    # Page Margins specified in template:
    # Left Margin: 3.81 cm (1.5 in)
    # Right Margin: 2.54 cm (1.0 in)
    # Top Margin: 2.54 cm (1.0 in)
    # Bottom Margin: 2.54 cm (1.0 in)
    section.left_margin = Cm(3.81)
    section.right_margin = Cm(2.54)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)

    # Set base Normal style font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.line_spacing = 1.5
    style.paragraph_format.space_after = Pt(6)

    # Setup footer page numbering
    setup_footer_page_number(section)

    return doc


def setup_footer_page_number(section):
    """Adds standard Arabic page number centered at the bottom of the page."""
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.font.name = 'Times New Roman'
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(50, 50, 50)

    fld1 = OxmlElement('w:fldChar')
    fld1.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = 'PAGE'
    fld2 = OxmlElement('w:fldChar')
    fld2.set(qn('w:fldCharType'), 'separate')
    fld3 = OxmlElement('w:fldChar')
    fld3.set(qn('w:fldCharType'), 'end')
    run._r.append(fld1)
    run._r.append(instr)
    run._r.append(fld2)
    run._r.append(fld3)


def add_p(doc, text: str, bold: bool = False, italic: bool = False,
          align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before: float = 0, space_after: float = 6):
    """Add standard justified 12pt Times New Roman paragraph with 1.5 line spacing."""
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)

    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    run.bold = bold
    run.italic = italic
    return p


def add_chapter_heading(doc, num: int, title: str):
    """
    Template Spec:
    Chapter Heading – Font Size: 16, Bold, Times New Roman, Centered
    """
    p1 = doc.add_paragraph()
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.space_before = Pt(24)
    p1.paragraph_format.space_after = Pt(6)
    r1 = p1.add_run(f"CHAPTER {num}")
    r1.font.name = 'Times New Roman'
    r1.font.size = Pt(16)
    r1.bold = True

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after = Pt(18)
    r2 = p2.add_run(title.upper())
    r2.font.name = 'Times New Roman'
    r2.font.size = Pt(16)
    r2.bold = True


def add_section_heading(doc, num_str: str, title: str):
    """
    Template Spec:
    Section Heading – Font Size: 14, Times New Roman, CAPS, Bold
    """
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True

    run = p.add_run(f"{num_str} {title.upper()}")
    run.font.name = 'Times New Roman'
    run.font.size = Pt(14)
    run.bold = True


def add_subsection_heading(doc, num_str: str, title: str):
    """
    Template Spec:
    Subsection Heading – Font Size: 12, Times New Roman, CAPS, Bold
    """
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True

    run = p.add_run(f"{num_str} {title.upper()}")
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    run.bold = True


def add_equation(doc, eq_str: str, eq_num: str = None, description: str = None):
    """Formats mathematical equations with indentation and equation numbering."""
    tbl = doc.add_table(rows=1, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False

    cell_eq, cell_num = tbl.rows[0].cells
    cell_eq.width = Inches(5.8)
    cell_num.width = Inches(1.0)

    p_eq = cell_eq.paragraphs[0]
    p_eq.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_eq.paragraph_format.space_after = Pt(2)
    r_eq = p_eq.add_run(eq_str)
    r_eq.font.name = 'Times New Roman'
    r_eq.font.size = Pt(11.5)
    r_eq.italic = True
    r_eq.bold = True

    p_num = cell_num.paragraphs[0]
    p_num.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_num.paragraph_format.space_after = Pt(2)
    if eq_num:
        r_num = p_num.add_run(f"({eq_num})")
        r_num.font.name = 'Times New Roman'
        r_num.font.size = Pt(11)
        r_num.bold = True

    # Remove borders
    for cell in (cell_eq, cell_num):
        tcPr = cell._element.get_or_add_tcPr()
        tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}>'
                              f'<w:top w:val="none"/>'
                              f'<w:left w:val="none"/>'
                              f'<w:bottom w:val="none"/>'
                              f'<w:right w:val="none"/>'
                              f'</w:tcBorders>')
        tcPr.append(tcBorders)

    if description:
        p_desc = doc.add_paragraph()
        p_desc.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_desc.paragraph_format.space_after = Pt(6)
        r_desc = p_desc.add_run(f"where: {description}")
        r_desc.font.name = 'Times New Roman'
        r_desc.font.size = Pt(10.5)
        r_desc.italic = True


def add_table_data(doc, table_num: str, title: str, headers: list, rows: list, col_widths: list = None):
    """
    Template Spec:
    Captions for tables can be in smaller fonts, but not smaller than 10pt.
    Adds a formally styled bordered table with shaded headers.
    """
    # Caption above table
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(10)
    p_cap.paragraph_format.space_after = Pt(4)
    p_cap.paragraph_format.keep_with_next = True

    r_bold = p_cap.add_run(f"Table {table_num}: ")
    r_bold.font.name = 'Times New Roman'
    r_bold.font.size = Pt(11)
    r_bold.bold = True

    r_title = p_cap.add_run(title)
    r_title.font.name = 'Times New Roman'
    r_title.font.size = Pt(11)

    table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Header Row
    hdr_cells = table.rows[0].cells
    for i, h_text in enumerate(headers):
        hdr_cells[i].text = h_text
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        for run in p.runs:
            run.font.name = 'Times New Roman'
            run.font.size = Pt(10.5)
            run.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)

        # Dark Navy Blue header background
        tcPr = hdr_cells[i]._element.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="1F3864"/>')
        tcPr.append(shd)

        # Padding
        mar = parse_xml(f'<w:tcMar {nsdecls("w")}>'
                        f'<w:top w:w="120" w:type="dxa"/>'
                        f'<w:bottom w:w="120" w:type="dxa"/>'
                        f'<w:left w:w="140" w:type="dxa"/>'
                        f'<w:right w:w="140" w:type="dxa"/>'
                        f'</w:tcMar>')
        tcPr.append(mar)

    # Data Rows
    for r_idx, row_data in enumerate(rows):
        row_cells = table.rows[r_idx + 1].cells
        bg_color = "F2F4F7" if (r_idx % 2 == 1) else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            row_cells[c_idx].text = str(val)
            p = row_cells[c_idx].paragraphs[0]
            # Center numbers or short tags, left align long text
            if len(str(val)) < 15 and not any(ch.isalpha() and len(str(val).split()) > 1 for ch in str(val)):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            for run in p.runs:
                run.font.name = 'Times New Roman'
                run.font.size = Pt(10)

            tcPr = row_cells[c_idx]._element.get_or_add_tcPr()
            shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg_color}"/>')
            tcPr.append(shd)
            mar = parse_xml(f'<w:tcMar {nsdecls("w")}>'
                            f'<w:top w:w="100" w:type="dxa"/>'
                            f'<w:bottom w:w="100" w:type="dxa"/>'
                            f'<w:left w:w="140" w:type="dxa"/>'
                            f'<w:right w:w="140" w:type="dxa"/>'
                            f'</w:tcMar>')
            tcPr.append(mar)

    # Apply Column Widths if provided
    if col_widths and len(col_widths) == len(headers):
        for row in table.rows:
            for idx, w in enumerate(col_widths):
                row.cells[idx].width = Inches(w)

    # Set Borders for the table
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(f'<w:tblBorders {nsdecls("w")}>'
                            f'<w:top w:val="single" w:sz="6" w:space="0" w:color="A6B9D0"/>'
                            f'<w:left w:val="none"/>'
                            f'<w:bottom w:val="single" w:sz="8" w:space="0" w:color="1F3864"/>'
                            f'<w:right w:val="none"/>'
                            f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="D9E1F2"/>'
                            f'<w:insideV w:val="none"/>'
                            f'</w:tblBorders>')
        tblPr[0].append(borders)

    # Space after table
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_after = Pt(6)


def add_figure_placeholder(doc, fig_num: str, title: str, description: str,
                           box_text: str = "[INSERT FIGURE HERE]"):
    """
    Template Spec:
    Captions for figures can be in smaller fonts, but not smaller than 10pt.
    Creates a distinct shaded schematic/figure placeholder box with comprehensive caption.
    """
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.rows[0].cells[0]
    cell.width = Inches(6.5)

    tcPr = cell._element.get_or_add_tcPr()
    # Light gray-blue box background
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F4F6F9"/>')
    tcPr.append(shd)
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}>'
                        f'<w:top w:val="single" w:sz="8" w:space="0" w:color="8FAADC"/>'
                        f'<w:left w:val="single" w:sz="8" w:space="0" w:color="8FAADC"/>'
                        f'<w:bottom w:val="single" w:sz="8" w:space="0" w:color="8FAADC"/>'
                        f'<w:right w:val="single" w:sz="8" w:space="0" w:color="8FAADC"/>'
                        f'</w:tcBorders>')
    tcPr.append(borders)
    mar = parse_xml(f'<w:tcMar {nsdecls("w")}>'
                    f'<w:top w:w="180" w:type="dxa"/>'
                    f'<w:bottom w:w="180" w:type="dxa"/>'
                    f'<w:left w:w="220" w:type="dxa"/>'
                    f'<w:right w:w="220" w:type="dxa"/>'
                    f'</w:tcMar>')
    tcPr.append(mar)

    p1 = cell.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.space_before = Pt(8)
    p1.paragraph_format.space_after = Pt(4)
    r1 = p1.add_run(f"─── {box_text} ───")
    r1.font.name = 'Times New Roman'
    r1.font.size = Pt(11)
    r1.bold = True
    r1.font.color.rgb = RGBColor(31, 56, 100)

    p2 = cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p2.paragraph_format.space_before = Pt(2)
    p2.paragraph_format.space_after = Pt(6)
    p2.paragraph_format.line_spacing = 1.15
    r2 = p2.add_run(description)
    r2.font.name = 'Times New Roman'
    r2.font.size = Pt(10)
    r2.italic = True
    r2.font.color.rgb = RGBColor(60, 60, 60)

    # Caption beneath figure
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(4)
    p_cap.paragraph_format.space_after = Pt(10)
    p_cap.paragraph_format.keep_with_next = False

    r_bold = p_cap.add_run(f"Figure {fig_num}: ")
    r_bold.font.name = 'Times New Roman'
    r_bold.font.size = Pt(10.5)
    r_bold.bold = True

    r_title = p_cap.add_run(title)
    r_title.font.name = 'Times New Roman'
    r_title.font.size = Pt(10.5)


def add_code_snippet(doc, code_str: str, caption: str = None):
    """Render monospace source code block in neat box."""
    if caption:
        p_cap = doc.add_paragraph()
        p_cap.paragraph_format.space_before = Pt(6)
        p_cap.paragraph_format.space_after = Pt(2)
        r = p_cap.add_run(f"Listing: {caption}")
        r.font.name = 'Times New Roman'
        r.font.size = Pt(10.5)
        r.bold = True

    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.rows[0].cells[0]
    cell.width = Inches(6.5)

    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F8F9FA"/>')
    tcPr.append(shd)
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}>'
                        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="D0D5DD"/>'
                        f'<w:left w:val="single" w:sz="6" w:space="0" w:color="1F3864"/>'
                        f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="D0D5DD"/>'
                        f'<w:right w:val="single" w:sz="4" w:space="0" w:color="D0D5DD"/>'
                        f'</w:tcBorders>')
    tcPr.append(borders)

    p = cell.paragraphs[0]
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(code_str)
    r.font.name = 'Courier New'
    r.font.size = Pt(9.0)
    r.font.color.rgb = RGBColor(20, 20, 20)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)
